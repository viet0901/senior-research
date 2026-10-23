"""Prepare one three-digit industry AI-exposure file for the thesis.

This is the only heavy ACS preparation step. It reads the preserved national
2015-2019 ACS PUMS person file, applies the documented worker sample, links
occupations to the Eloundou et al. score, and aggregates the score to QWI
three-digit industries with person weights.

Run from the project root:

    python code/prepare_ai_exposure.py

Analysis, figures, the employment trend, panel construction, and regressions
belong in the research notebook rather than in this file.
"""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw"
OUTPUT = ROOT / "data/interim/ai_exposure"

PUMS_FILE = RAW / "acs_pums/2015_2019_national/csv_pus.zip"
AI_SCORE_FILE = RAW / "ai_exposure/occ_level.csv"
ONET_SOC_CROSSWALK = RAW / "acs_pums/crosswalks/onet_2019_to_soc2018_crosswalk.csv"
ACS_OCCUPATION_CROSSWALK = (
    RAW / "acs_pums/crosswalks/2018-ACS-PUMS-and-2018-SIPP-Public-Use-Occupation-Code-List.xlsx"
)
INDUSTRY_CROSSWALK = RAW / "acs_pums/crosswalks/2022-Census-Industry-Code-List-with-Crosswalk.xlsx"
QWI_VALIDATION = ROOT / "data/interim/qwi/qwi_research_ind3_validation.json"

PUMS_COLUMNS = ["ESR", "COW", "PWGTP", "INDP", "OCCP"]
EMPLOYED = {"1", "2"}
PRIVATE_EMPLOYEE = {"1", "2"}


def require_files(paths: list[Path]) -> None:
    """Stop with a useful message if a required input is missing."""
    missing = [str(path.relative_to(ROOT)) for path in paths if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing research input(s): " + ", ".join(missing))


def normal_code(value: object, width: int) -> str:
    """Return a zero-padded Census code while preserving missing values."""
    if value is None or pd.isna(value):
        return "<missing>"
    code = str(value).strip()
    if not code or code.lower() in {"nan", "none", "bbbb"}:
        return "<missing>"
    return code.zfill(width)


def read_qwi_industries() -> tuple[set[str], dict[str, str]]:
    """Read the three-digit industry codes from the QWI run's saved schema."""
    report = json.loads(QWI_VALIDATION.read_text(encoding="utf-8"))
    labels_file = ROOT / report["raw_directory"] / "schema/label_industry.csv"
    labels = pd.read_csv(labels_file, dtype="string")
    labels = labels.loc[labels["ind_level"].eq("3"), ["industry", "label"]]
    names = dict(labels.itertuples(index=False, name=None))
    return set(names), names


def read_acs_occupation_codes() -> dict[str, dict[str, object]]:
    """Read ACS occupation codes and their official 2018 SOC members."""
    workbook = openpyxl.load_workbook(
        ACS_OCCUPATION_CROSSWALK, read_only=True, data_only=True
    )
    sheet = workbook["ACS"]
    occupations: dict[str, dict[str, object]] = {}
    current: str | None = None
    in_combines_block = False

    for raw_code, raw_soc, raw_title in sheet.iter_rows(
        min_row=11, min_col=1, max_col=3, values_only=True
    ):
        code = str(raw_code).strip() if raw_code is not None else ""
        soc = str(raw_soc).strip() if raw_soc is not None else ""
        title = str(raw_title).strip() if raw_title is not None else ""

        if re.fullmatch(r"\d{4}", code):
            current = code
            occupations[current] = {"title": title, "soc_header": soc, "members": []}
            in_combines_block = False
        elif current and soc == "Combines:":
            in_combines_block = True
        elif current and in_combines_block and re.fullmatch(r"\d{2}-\d{4}", soc):
            occupations[current]["members"].append(soc.replace("-", ""))

    for detail in occupations.values():
        header = str(detail["soc_header"])
        if not detail["members"] and re.fullmatch(r"\d{2}-\d{4}", header):
            detail["members"] = [header.replace("-", "")]

    if not occupations:
        raise RuntimeError("No occupation codes were found in the ACS code list.")
    return occupations


def build_occupation_scores() -> pd.DataFrame:
    """Assign a score only when all official SOC members have one score."""
    scores = pd.read_csv(AI_SCORE_FILE, dtype={"O*NET-SOC Code": "string"})
    scores["dv_rating_beta"] = pd.to_numeric(scores["dv_rating_beta"], errors="raise")
    if scores["O*NET-SOC Code"].duplicated().any():
        raise RuntimeError("The AI-score source contains duplicate occupation codes.")
    if scores["dv_rating_beta"].isna().any() or not scores["dv_rating_beta"].between(0, 1).all():
        raise RuntimeError("AI exposure scores must be complete and between zero and one.")

    crosswalk = pd.read_csv(ONET_SOC_CROSSWALK, dtype="string")
    linked = scores.merge(
        crosswalk[["O*NET-SOC 2019 Code", "2018 SOC Code"]],
        left_on="O*NET-SOC Code",
        right_on="O*NET-SOC 2019 Code",
        how="left",
        validate="one_to_one",
    )
    linked["soc2018"] = linked["2018 SOC Code"].str.replace("-", "", regex=False)
    score_by_soc = (
        linked.dropna(subset=["soc2018"])
        .groupby("soc2018")["dv_rating_beta"]
        .agg(lambda values: sorted(set(float(value) for value in values)))
        .to_dict()
    )

    rows = []
    for occp, detail in read_acs_occupation_codes().items():
        members = list(detail["members"])
        member_scores = [score_by_soc.get(soc, []) for soc in members]
        distinct_scores = sorted({score for values in member_scores for score in values})

        if not members:
            status = "no_exact_soc_members"
            score = None
        elif any(soc.endswith("0000") for soc in members):
            status = "broad_soc_member"
            score = None
        elif any(not values for values in member_scores):
            status = "unscored_soc_member"
            score = None
        elif len(distinct_scores) == 1:
            status = "matched"
            score = distinct_scores[0]
        else:
            status = "multiple_scores"
            score = None

        rows.append(
            {
                "OCCP": occp,
                "occupation_title": detail["title"],
                "occupation_match_status": status,
                "ai_exposure": score,
            }
        )

    return pd.DataFrame(rows)


def build_industry_map(qwi_codes: set[str]) -> dict[str, str]:
    """Map a 2017 ACS industry only when all 2022 targets share one prefix."""
    workbook = openpyxl.load_workbook(INDUSTRY_CROSSWALK, read_only=True, data_only=True)
    sheet = workbook["2007 to 2022 Industry Crosswalk"]
    mapping: dict[str, str] = {}

    for row in sheet.iter_rows(min_row=8, values_only=True):
        old_code, new_codes = row[8], row[13]
        if old_code is None or new_codes is None:
            continue
        old = str(old_code).strip().zfill(4)
        options = [item.strip() for item in str(new_codes).split(",")]
        if not re.fullmatch(r"\d{4}", old):
            continue
        if not options or any(not re.fullmatch(r"\d{3,6}", item) for item in options):
            continue
        targets = {item[:3] for item in options}
        if len(targets) == 1 and next(iter(targets)) in qwi_codes:
            mapping[old] = next(iter(targets))

    return mapping


def aggregate_pums(
    occupation_scores: pd.DataFrame,
    industry_map: dict[str, str],
) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, int]]:
    """Stream PUMS and retain only small weighted summaries."""
    score_map = occupation_scores.set_index("OCCP")["ai_exposure"].to_dict()
    status_map = occupation_scores.set_index("OCCP")["occupation_match_status"].to_dict()
    totals = Counter()
    industries: dict[str, Counter] = defaultdict(Counter)
    occupations: dict[str, Counter] = defaultdict(Counter)

    with ZipFile(PUMS_FILE) as archive:
        members = sorted(
            name
            for name in archive.namelist()
            if Path(name).name.lower().startswith("psam_pus") and name.lower().endswith(".csv")
        )
        if len(members) != 4:
            raise RuntimeError(f"Expected four national PUMS person files, found {len(members)}.")

        for member in members:
            print(f"Reading {Path(member).name}")
            with archive.open(member) as source:
                for chunk in pd.read_csv(
                    source, usecols=PUMS_COLUMNS, dtype="string", chunksize=100_000
                ):
                    weights = pd.to_numeric(chunk["PWGTP"], errors="raise").astype("int64")
                    totals["all_person_records"] += len(chunk)
                    eligible = chunk["ESR"].isin(EMPLOYED) & chunk["COW"].isin(PRIVATE_EMPLOYEE)
                    workers = chunk.loc[eligible].copy()
                    worker_weights = weights.loc[eligible]
                    totals["eligible_worker_records"] += len(workers)
                    totals["eligible_worker_weight"] += int(worker_weights.sum())

                    indp = workers["INDP"].map(lambda value: normal_code(value, 4))
                    occp = workers["OCCP"].map(lambda value: normal_code(value, 4))
                    target = indp.map(industry_map)
                    exposure = occp.map(score_map)

                    totals["industry_mapped_weight"] += int(worker_weights[target.notna()].sum())
                    totals["occupation_scored_weight"] += int(worker_weights[exposure.notna()].sum())
                    joint = target.notna() & exposure.notna()
                    totals["jointly_matched_weight"] += int(worker_weights[joint].sum())

                    industry_chunk = pd.DataFrame(
                        {
                            "industry": target,
                            "weight": worker_weights.to_numpy(),
                            "weighted_score": worker_weights.to_numpy() * exposure.fillna(0).to_numpy(),
                            "matched_weight": worker_weights.to_numpy() * exposure.notna().to_numpy(),
                        }
                    ).dropna(subset=["industry"])
                    for industry, values in industry_chunk.groupby("industry"):
                        industries[str(industry)]["eligible_person_weight"] += int(values["weight"].sum())
                        industries[str(industry)]["matched_person_weight"] += int(values["matched_weight"].sum())
                        industries[str(industry)]["weighted_score"] += float(values["weighted_score"].sum())

                    occupation_chunk = pd.DataFrame(
                        {"OCCP": occp, "weight": worker_weights.to_numpy()}
                    ).groupby("OCCP", as_index=False)["weight"].sum()
                    for code, weight in occupation_chunk.itertuples(index=False, name=None):
                        occupations[str(code)]["person_weight"] += int(weight)

    if totals["all_person_records"] != 15_947_624:
        raise RuntimeError("The PUMS record count does not match the Census national control.")

    industry_rows = []
    for industry, values in industries.items():
        matched = int(values["matched_person_weight"])
        eligible = int(values["eligible_person_weight"])
        industry_rows.append(
            {
                "industry": industry,
                "eligible_person_weight": eligible,
                "matched_person_weight": matched,
                "ai_exposure_coverage": matched / eligible if eligible else None,
                "ai_exposure": float(values["weighted_score"] / matched) if matched else None,
            }
        )

    occupation_rows = []
    for code, values in occupations.items():
        occupation_rows.append(
            {
                "OCCP": code,
                "person_weight": int(values["person_weight"]),
                "occupation_match_status": status_map.get(code, "not_in_crosswalk"),
                "ai_exposure": score_map.get(code),
            }
        )

    return pd.DataFrame(industry_rows), pd.DataFrame(occupation_rows), dict(totals)


def run() -> dict[str, object]:
    """Create the small exposure and validation files used by the notebook."""
    require_files(
        [
            PUMS_FILE,
            AI_SCORE_FILE,
            ONET_SOC_CROSSWALK,
            ACS_OCCUPATION_CROSSWALK,
            INDUSTRY_CROSSWALK,
            QWI_VALIDATION,
        ]
    )
    qwi_codes, qwi_names = read_qwi_industries()
    occupation_scores = build_occupation_scores()
    industry_map = build_industry_map(qwi_codes)
    exposure, occupation_coverage, totals = aggregate_pums(occupation_scores, industry_map)

    exposure = pd.DataFrame({"industry": sorted(qwi_codes)}).merge(
        exposure, on="industry", how="left", validate="one_to_one"
    )
    exposure["industry_name"] = exposure["industry"].map(qwi_names)
    exposure = exposure[
        [
            "industry",
            "industry_name",
            "ai_exposure",
            "ai_exposure_coverage",
            "eligible_person_weight",
            "matched_person_weight",
        ]
    ]

    OUTPUT.mkdir(parents=True, exist_ok=True)
    exposure.to_csv(OUTPUT / "industry_ai_exposure_3digit.csv", index=False)
    occupation_coverage.sort_values("person_weight", ascending=False).to_csv(
        OUTPUT / "occupation_coverage.csv", index=False
    )

    eligible_weight = totals["eligible_worker_weight"]
    report: dict[str, object] = {
        "status": "Prepared input for notebook analysis; no panel or regression was created.",
        "worker_sample": "ESR 1/2, COW 1/2, weighted by PWGTP",
        "industry_level": "three-digit QWI industry",
        "occupation_rule": (
            "An ACS occupation receives a score only when every detailed SOC member "
            "has the same available Eloundou beta score."
        ),
        "counts": totals,
        "coverage": {
            "unique_industry_share": totals["industry_mapped_weight"] / eligible_weight,
            "scored_occupation_share": totals["occupation_scored_weight"] / eligible_weight,
            "joint_match_share": totals["jointly_matched_weight"] / eligible_weight,
            "qwi_industries": len(qwi_codes),
            "industries_with_exposure": int(exposure["ai_exposure"].notna().sum()),
        },
        "outputs": [
            "data/interim/ai_exposure/industry_ai_exposure_3digit.csv",
            "data/interim/ai_exposure/occupation_coverage.csv",
        ],
    }
    (OUTPUT / "ai_exposure_validation.json").write_text(
        json.dumps(report, indent=2) + "\n", encoding="utf-8"
    )
    return report


if __name__ == "__main__":
    result = run()
    print(json.dumps(result["coverage"], indent=2))
