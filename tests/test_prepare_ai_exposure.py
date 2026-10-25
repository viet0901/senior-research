"""Unit tests for conservative industry and occupation matching."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest

import prepare_ai_exposure


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("170", "0170"),
        (17, "0017"),
        (" 0170 ", "0170"),
        (None, "<missing>"),
        (pd.NA, "<missing>"),
        ("bbbb", "<missing>"),
    ],
)
def test_normal_code_preserves_codes_and_missing_values(value: object, expected: str) -> None:
    assert prepare_ai_exposure.normal_code(value, 4) == expected


def test_build_occupation_scores_keeps_only_unambiguous_scores(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    score_file = tmp_path / "scores.csv"
    crosswalk_file = tmp_path / "crosswalk.csv"
    pd.DataFrame(
        {
            "O*NET-SOC Code": ["11-1011.00", "13-1111.00"],
            "dv_rating_beta": [0.8, 0.2],
        }
    ).to_csv(score_file, index=False)
    pd.DataFrame(
        {
            "O*NET-SOC 2019 Code": ["11-1011.00", "13-1111.00"],
            "2018 SOC Code": ["11-1011", "13-1111"],
        }
    ).to_csv(crosswalk_file, index=False)

    occupations = {
        "0010": {"title": "Unique", "soc_header": "", "members": ["111011"]},
        "0020": {
            "title": "Different scores",
            "soc_header": "",
            "members": ["111011", "131111"],
        },
        "0030": {"title": "No detail", "soc_header": "", "members": []},
        "0040": {"title": "Broad", "soc_header": "", "members": ["110000"]},
        "0050": {"title": "Unscored", "soc_header": "", "members": ["151252"]},
    }
    monkeypatch.setattr(prepare_ai_exposure, "AI_SCORE_FILE", score_file)
    monkeypatch.setattr(prepare_ai_exposure, "ONET_SOC_CROSSWALK", crosswalk_file)
    monkeypatch.setattr(prepare_ai_exposure, "read_acs_occupation_codes", lambda: occupations)

    result = prepare_ai_exposure.build_occupation_scores().set_index("OCCP")

    assert result.loc["0010", "occupation_match_status"] == "matched"
    assert result.loc["0010", "ai_exposure"] == pytest.approx(0.8)
    assert result.loc["0020", "occupation_match_status"] == "multiple_scores"
    assert result.loc["0030", "occupation_match_status"] == "no_exact_soc_members"
    assert result.loc["0040", "occupation_match_status"] == "broad_soc_member"
    assert result.loc["0050", "occupation_match_status"] == "unscored_soc_member"
    assert result.loc[["0020", "0030", "0040", "0050"], "ai_exposure"].isna().all()


def test_build_industry_map_accepts_only_one_valid_three_digit_target(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook_path = tmp_path / "industry_crosswalk.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "2007 to 2022 Industry Crosswalk"
    sheet.cell(row=8, column=9, value="0170")
    sheet.cell(row=8, column=14, value="111")
    sheet.cell(row=9, column=9, value="0180")
    sheet.cell(row=9, column=14, value="111, 112")
    sheet.cell(row=10, column=9, value="0190")
    sheet.cell(row=10, column=14, value="999")
    sheet.cell(row=11, column=9, value="0270")
    sheet.cell(row=11, column=14, value="112, 1123")
    workbook.save(workbook_path)
    monkeypatch.setattr(prepare_ai_exposure, "INDUSTRY_CROSSWALK", workbook_path)

    result = prepare_ai_exposure.build_industry_map({"111", "112"})

    assert result == {"0170": "111", "0270": "112"}


def test_require_files_names_missing_inputs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(prepare_ai_exposure, "ROOT", tmp_path)
    existing = tmp_path / "existing.csv"
    existing.write_text("present\n", encoding="utf-8")
    missing = tmp_path / "missing.csv"

    with pytest.raises(FileNotFoundError, match="missing.csv"):
        prepare_ai_exposure.require_files([existing, missing])
